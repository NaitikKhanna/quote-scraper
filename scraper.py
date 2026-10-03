import time

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def scrape_page(url):
    try:
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        response.encoding = 'utf-8'
    except requests.RequestException as e:
        print(f"Failed to fetch {url}: {e}")
        return []

    soup = BeautifulSoup(response.text, 'html.parser')
    products = []

    for item in soup.find_all("div", class_="quote"):
        quote_tag = item.find("span", class_="text")
        author_tag = item.find("small", class_="author")
        about_tag = item.find("a", href=True, string="(about)")

        if not quote_tag or not author_tag or not about_tag:
            print("Skipped one item — missing quote, author, or link")
            continue

        products.append({
            'quote': quote_tag.text.strip(),
            'author': author_tag.text.strip(),
            'about': "https://quotes.toscrape.com" + about_tag['href']
        })

    return products


def scrape_all_pages(base_url, max_pages=1):
    all_products = []

    for page_num in range(1, max_pages + 1):
        if page_num == 1:
            url = base_url
        else:
            url = f"{base_url}page/{page_num}/"

        print(f"Scraping page {page_num}...")
        products = scrape_page(url)

        if not products:
            break

        all_products.extend(products)
        time.sleep(1)

    return all_products


def clean_data(data):
    df = pd.DataFrame(data)

    if df.empty:
        print("No data scraped.")
        return df

    df = df.drop_duplicates(subset=['quote'])

    return df


def save_pretty_excel(df, filename):
    df.to_excel(filename, index=False, sheet_name='Quotes')

    wb = load_workbook(filename)
    ws = wb['Quotes']

    header_fill = PatternFill(start_color="1C2B3A", end_color="1C2B3A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value)) for cell in column_cells)
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = min(max_length + 4, 60)

    wb.save(filename)
    print(f"Saved formatted Excel file: {filename}")


if __name__ == "__main__":
    data = scrape_all_pages(
        "https://quotes.toscrape.com/",
        max_pages=3
    )

    df = clean_data(data)

    if not df.empty:
        save_pretty_excel(df, "scraped_quotes.xlsx")

    print(df)
    print(f"\nTotal quotes scraped: {len(df)}")